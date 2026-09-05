"""Safe local storage and bounded text extraction for ordinary email attachments."""
import hashlib
import re
from dataclasses import dataclass, field
from pathlib import Path
from typing import Any

from email_agent.domain.models import EmailAttachment


@dataclass(frozen=True)
class StoredAttachment:
    attachment_id: str
    filename: str
    content_type: str
    source: str
    path: str
    size_bytes: int
    extracted_text: str = ""
    extraction_status: str = "not_extracted"
    extraction_error: str = ""
    metadata: dict[str, Any] = field(default_factory=dict)


class AttachmentStorage:
    PDF_TYPES = {"application/pdf"}
    DOCX_TYPES = {
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
    }
    XLSX_TYPES = {
        "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
    }
    TEXT_TYPES = {"text/csv", "text/plain"}

    def __init__(self, root: Path, config: dict | None = None):
        self.root = Path(root)
        self.config = config or {}
        self.max_attachments_per_email = int(self.config.get("max_attachments_per_email", 8))
        self.max_attachment_bytes = int(self.config.get("max_attachment_bytes", 10 * 1024 * 1024))
        self.max_extracted_chars = int(self.config.get("max_extracted_chars", 12000))

    def save(
        self, message_id: str, attachments: list[EmailAttachment]
    ) -> tuple[list[StoredAttachment], list[str]]:
        message_dir = self._message_dir(message_id)
        message_dir.mkdir(parents=True, exist_ok=True)
        stored, errors = [], []
        for index, item in enumerate(attachments, 1):
            if index > self.max_attachments_per_email:
                errors.append(f"{item.filename} 超过单封邮件附件数量限制")
                continue
            size = len(item.data)
            if size > self.max_attachment_bytes:
                errors.append(f"{item.filename} 超过附件大小限制")
                continue

            target = self._unique_path(message_dir, item.filename, item.attachment_id)
            target.write_bytes(item.data)
            extracted_text, status, extraction_error = self._extract(item, target)
            stored.append(
                StoredAttachment(
                    attachment_id=item.attachment_id,
                    filename=target.name,
                    content_type=item.content_type,
                    source=item.source,
                    path=str(target),
                    size_bytes=target.stat().st_size,
                    extracted_text=extracted_text,
                    extraction_status=status,
                    extraction_error=extraction_error,
                    metadata=self._safe_metadata(item.metadata),
                )
            )
        return stored, errors

    @staticmethod
    def manifest(stored: list[StoredAttachment]) -> list[dict]:
        return [
            {
                "attachment_id": item.attachment_id,
                "filename": item.filename,
                "content_type": item.content_type,
                "source": item.source,
                "path": item.path,
                "size_bytes": item.size_bytes,
                "extracted_text": item.extracted_text,
                "extraction_status": item.extraction_status,
                "extraction_error": item.extraction_error,
                "metadata": item.metadata,
            }
            for item in stored
        ]

    def _message_dir(self, message_id: str) -> Path:
        digest = hashlib.sha256(str(message_id or "unknown").encode("utf-8")).hexdigest()[:32]
        directory = (self.root / digest).resolve()
        root = self.root.resolve()
        if directory.parent != root:
            raise ValueError("附件目录路径越界")
        return directory

    def _unique_path(self, directory: Path, filename: str, attachment_id: str) -> Path:
        clean = self._safe_filename(filename)
        stem = Path(clean).stem[:80] or "attachment"
        suffix = Path(clean).suffix.lower()[:10] or ".bin"
        item_digest = hashlib.sha256(str(attachment_id).encode("utf-8")).hexdigest()[:8]
        target = (directory / f"{stem}-{item_digest}{suffix}").resolve()
        if directory.resolve() not in target.parents:
            raise ValueError("附件文件路径越界")
        return target

    @staticmethod
    def _safe_filename(filename: str) -> str:
        name = Path(str(filename or "")).name.strip()
        name = re.sub(r'[<>:"/\\|?*\x00-\x1f]', "_", name).strip(" ._")
        return name or "attachment.bin"

    def _extract(self, item: EmailAttachment, path: Path) -> tuple[str, str, str]:
        content_type = item.content_type.lower().split(";", 1)[0].strip()
        try:
            if content_type in self.PDF_TYPES:
                text = self._extract_pdf(path, self.max_extracted_chars)
            elif content_type in self.DOCX_TYPES:
                text = self._extract_docx(path, self.max_extracted_chars)
            elif content_type in self.XLSX_TYPES:
                text = self._extract_xlsx(path, self.max_extracted_chars)
            elif content_type in self.TEXT_TYPES or path.suffix.lower() in {".csv", ".txt"}:
                text = path.read_bytes()[: self.max_extracted_chars * 4].decode(
                    "utf-8", errors="replace"
                )
            else:
                return "", "metadata_only", ""
        except Exception as exc:
            return "", "extraction_failed", str(exc)
        return text[: self.max_extracted_chars], "extracted", ""

    @staticmethod
    def _extract_pdf(path: Path, max_chars: int) -> str:
        from pypdf import PdfReader

        reader = PdfReader(str(path))
        lines = []
        total = 0
        for page in reader.pages:
            if total >= max_chars:
                break
            text = page.extract_text() or ""
            if not text:
                continue
            lines.append(text[: max_chars - total])
            total += len(lines[-1])
        return "\n".join(lines)

    @staticmethod
    def _safe_metadata(value) -> dict[str, Any]:
        if not isinstance(value, dict):
            return {}

        def clean(item):
            if isinstance(item, (bytes, bytearray, memoryview)):
                return None
            if isinstance(item, dict):
                cleaned = {}
                for key, child in item.items():
                    if str(key).lower() in {"data", "payload", "content", "raw"}:
                        continue
                    cleaned_child = clean(child)
                    if cleaned_child is not None:
                        cleaned[str(key)[:80]] = cleaned_child
                return cleaned
            if isinstance(item, (list, tuple)):
                cleaned_items = [clean(child) for child in item[:20]]
                return [child for child in cleaned_items if child is not None]
            if isinstance(item, (str, int, float, bool)) or item is None:
                return item[:500] if isinstance(item, str) else item
            return str(item)[:200]

        return clean(value) or {}

    @staticmethod
    def _extract_docx(path: Path, max_chars: int) -> str:
        from docx import Document

        document = Document(str(path))
        lines = []
        total = 0
        for paragraph in document.paragraphs:
            text = paragraph.text
            if not text:
                continue
            remaining = max_chars - total
            if remaining <= 0:
                break
            lines.append(text[:remaining])
            total += len(lines[-1])
        return "\n".join(lines)

    @staticmethod
    def _extract_xlsx(path: Path, max_chars: int) -> str:
        from openpyxl import load_workbook

        workbook = load_workbook(str(path), read_only=True, data_only=True)
        try:
            rows = []
            total = 0
            for worksheet in workbook.worksheets:
                if total >= max_chars:
                    break
                header = f"[{worksheet.title}]"
                rows.append(header[: max_chars - total])
                total += len(rows[-1])
                for row in worksheet.iter_rows(values_only=True):
                    if total >= max_chars:
                        break
                    line = ",".join("" if value is None else str(value) for value in row)
                    rows.append(line[: max_chars - total])
                    total += len(rows[-1])
            return "\n".join(rows)
        finally:
            workbook.close()


def format_attachment_context(stored: list[StoredAttachment], max_chars: int) -> str:
    if max_chars <= 0:
        return ""
    parts = []
    for item in stored:
        if not item.extracted_text:
            continue
        parts.append(f"{item.filename}: {item.extracted_text}")
    return "\n".join(parts)[:max_chars]
